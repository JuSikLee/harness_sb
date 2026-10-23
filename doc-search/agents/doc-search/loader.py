"""사내 문서(.docx/.xlsx/.pdf) → 섹션 청크 로더. 표는 행 단위 청크 + 헤더 prepend."""
import hashlib
import json
import re
import sys
from pathlib import Path

import docx
import openpyxl
import pdfplumber
from docx.table import Table
from docx.text.paragraph import Paragraph

DOC_ROOT = Path("/home/prime/harness_sb/harness-book-samples/doc-search")
MAX_SECTION = 600
CHUNK_LIMIT = 300
SUFFIXES = {".docx", ".xlsx", ".pdf"}
PDF_HEADING = re.compile(r"^\d+\.\s+\S")
SENTENCE_END = re.compile(r"(?<=[.!?。])\s+|\n+")


def file_sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def list_files(root=DOC_ROOT):
    return sorted(p for p in Path(root).rglob("*") if p.suffix.lower() in SUFFIXES)


def corpus_hash(root=DOC_ROOT):
    h = hashlib.sha256()
    for p in list_files(root):
        h.update(f"{p.relative_to(root)}:{file_sha256(p)}".encode())
    return h.hexdigest()


def split_text(text):
    """600자 초과 섹션을 문단/문장 경계로 ≤300자 청크로 재분할."""
    if len(text) <= MAX_SECTION:
        return [text]
    pieces = []
    for sent in SENTENCE_END.split(text):
        sent = sent.strip()
        while len(sent) > CHUNK_LIMIT:  # 경계 없는 긴 문장은 강제 분할
            pieces.append(sent[:CHUNK_LIMIT])
            sent = sent[CHUNK_LIMIT:]
        if sent:
            pieces.append(sent)
    chunks, cur = [], ""
    for piece in pieces:
        if cur and len(cur) + len(piece) + 1 > CHUNK_LIMIT:
            chunks.append(cur)
            cur = piece
        else:
            cur = f"{cur} {piece}".strip()
    if cur:
        chunks.append(cur)
    return chunks


def clean(cell):
    return "" if cell is None else re.sub(r"\s+", " ", str(cell)).strip()


class Collector:
    def __init__(self, doc_name):
        self.doc = doc_name
        self.sections = []

    def add_text(self, heading_path, body):
        body = body.strip()
        if not body:
            return
        for piece in split_text(body):
            self.sections.append(self._make(heading_path, piece, "text"))

    def add_table(self, heading_path, rows):
        """첫 유효 행=헤더. 각 데이터 행을 헤더와 묶어 별도 청크로."""
        rows = [[clean(c) for c in r] for r in rows]
        rows = [r for r in rows if any(r)]
        if len(rows) < 2:
            for r in rows:
                self.sections.append(self._make(heading_path, " | ".join(c for c in r if c), "table_row"))
            return
        header, data = rows[0], rows[1:]
        for r in data:
            pairs = []
            for i, v in enumerate(r):
                if not v:
                    continue
                h = header[i] if i < len(header) and header[i] else ""
                pairs.append(f"{h}: {v}" if h else v)
            if pairs:
                self.sections.append(self._make(heading_path, " | ".join(pairs), "table_row"))

    def _make(self, heading_path, body, kind):
        path = " > ".join([self.doc] + [h for h in heading_path if h])
        return {"doc": self.doc, "path": path, "kind": kind, "body": body, "text": f"{path}\n{body}"}


def load_docx(path, rel):
    d = docx.Document(path)
    col = Collector(rel)
    stack, buf = {}, []

    def heading_path():
        return [stack[k] for k in sorted(stack)]

    def flush():
        col.add_text(heading_path(), "\n".join(buf))
        buf.clear()

    for child in d.element.body.iterchildren():
        if child.tag.endswith("}p"):
            p = Paragraph(child, d)
            text = p.text.strip()
            style = p.style.name if p.style is not None else ""
            if style.startswith("Heading") and text:
                flush()
                level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
                for k in [k for k in stack if k >= level]:
                    del stack[k]
                stack[level] = text
            elif text:
                buf.append(text)
        elif child.tag.endswith("}tbl"):
            flush()
            t = Table(child, d)
            rows = []
            for row in t.rows:
                cells, prev = [], None
                for c in row.cells:
                    if c._tc is prev:  # 병합 셀 중복 제거
                        continue
                    prev = c._tc
                    cells.append(c.text)
                rows.append(cells)
            col.add_table(heading_path(), rows)
    flush()
    return col.sections


def load_xlsx(path, rel):
    wb = openpyxl.load_workbook(path, data_only=True)
    col = Collector(rel)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(clean(c) for c in r)]
        title_parts = []
        # 단일 셀뿐인 앞쪽 행은 제목/설명 → 헤딩 경로로, 다음 다중 셀 행이 헤더
        while rows and sum(1 for c in rows[0] if clean(c)) < 2 and len(rows) > 1:
            title_parts.append(clean(next(c for c in rows[0] if clean(c))))
            rows = rows[1:]
        heading = [ws.title] + title_parts[:1]
        notes = title_parts[1:]
        if notes:
            col.add_text(heading, "\n".join(notes))
        col.add_table(heading, rows)
    return col.sections


def load_pdf(path, rel):
    col = Collector(rel)
    heading, buf = "", []

    def flush():
        col.add_text([heading], "\n".join(buf))
        buf.clear()

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            tables = page.find_tables()
            for t in tables:  # 표 우선
                flush()
                col.add_table([heading], t.extract())
            body = page
            for t in tables:
                body = body.outside_bbox(t.bbox)
            for line in (body.extract_text() or "").splitlines():
                line = line.strip()
                if not line:
                    continue
                if PDF_HEADING.match(line):
                    flush()
                    heading = line
                else:
                    buf.append(line)
        flush()
    return col.sections


LOADERS = {".docx": load_docx, ".xlsx": load_xlsx, ".pdf": load_pdf}


def load_all(root=DOC_ROOT):
    sections = []
    for p in list_files(root):
        sections.extend(LOADERS[p.suffix.lower()](p, str(p.relative_to(root))))
    for i, s in enumerate(sections):
        s["id"] = i
    return sections


if __name__ == "__main__":
    secs = load_all()
    print(f"섹션 수: {len(secs)}")
    if "--dump" in sys.argv:
        print(json.dumps(secs, ensure_ascii=False, indent=1))
