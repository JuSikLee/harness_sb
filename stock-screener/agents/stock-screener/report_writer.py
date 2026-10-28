"""결과 .docx 작성 (specs/architecture.md 섹션 6).

output/recommendations/{YYYYMMDD}_top10_{NN}.docx. 면책 문구(한국어 + 영어)를 첫 단락 · 마지막 단락에 넣고,
종목마다 상승 이유 · 산업 트렌드 · 하방 리스크 · 근거 출처 · 추천 등급을 표로 기록한다.
"""
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "output" / "recommendations"
FONT = "Malgun Gothic"

DISCLAIMER_KR = "본 자료는 정보 제공 목적이며 투자 권유가 아닙니다."
DISCLAIMER_EN = "This material is provided for informational purposes only and is not an offer or solicitation to invest."

# "매수/매도" 단정 표현 → "검토 후보"/"주의 필요" (순매수 · 매수세 같은 시장 용어는 건드리지 않는다)
_BUY = re.compile(r"(강력\s*)?매수\s*(하세요|하라|하십시오|추천|의견|권고|적기)|(지금|당장)\s*매수|\bstrong buy\b|\bbuy rating\b|\bbuy now\b|\byou should buy\b", re.I)
_SELL = re.compile(r"(강력\s*)?매도\s*(하세요|하라|하십시오|추천|의견|권고)|(지금|당장)\s*매도|\bstrong sell\b|\bsell rating\b|\bsell now\b|\byou should sell\b", re.I)


def mask_assertive(text: str) -> str:
    return _SELL.sub("주의 필요", _BUY.sub("검토 후보", text or ""))


def fmt_mcap(value: float, currency: str) -> str:
    if currency == "KRW":
        return f"{value / 1e12:,.1f}조원 (KRW)" if value >= 1e12 else f"{value / 1e8:,.0f}억원 (KRW)"
    return f"${value / 1e9:,.1f}B (USD)" if value >= 1e9 else f"${value / 1e6:,.0f}M (USD)"


def next_path(now: datetime | None = None) -> Path:
    now = now or datetime.now()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prefix = f"{now:%Y%m%d}_top10_"
    taken = [int(m.group(1)) for p in OUT_DIR.glob(prefix + "*.docx") if (m := re.search(r"_(\d+)\.docx$", p.name))]
    return OUT_DIR / f"{prefix}{max(taken, default=0) + 1:02d}.docx"


def _font(run, size=None, bold=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _para(doc, text, size=10, bold=False, align=None):
    p = doc.add_paragraph()
    _font(p.add_run(text), size, bold)
    if align:
        p.alignment = align
    return p


def _disclaimer(doc):
    p = doc.add_paragraph()
    _font(p.add_run(DISCLAIMER_KR), 10, True)
    p.add_run().add_break()
    _font(p.add_run(DISCLAIMER_EN), 9, False)


def _cell(cell, text, bold=False, size=9):
    cell.text = ""
    _font(cell.paragraphs[0].add_run(text), size, bold)


def _table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        _cell(t.rows[0].cells[i], h, True)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            _cell(cells[i], str(v))
    return t


def _cites(evidence) -> str:
    return "\n".join(f"[{a['title']}, {a['date']}, {a['media']}] {a['url']}" for a in evidence)


def write(picks: list[dict], meta: dict, path: Path | None = None) -> Path:
    """picks: 검증 · 등급 · 점수까지 끝난 종목 리스트. meta: request, stats, notes, excluded, tokens, elapsed."""
    path = path or next_path()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    _disclaimer(doc)  # 첫 단락
    h = doc.add_heading("주식 종목 추천 리포트 — 최근 5일 상승 종목 Top 10", level=1)
    for r in h.runs:
        _font(r, 16, True)
    _para(doc, f"작성일 {datetime.now():%Y-%m-%d %H:%M} · 요청: \"{meta['request']}\"")
    _para(doc, "데이터: 주가·거래량·시총 yfinance(한국 시총 pykrx 우선), 뉴스 Tavily, 공시 DART/SEC EDGAR. "
               "추천 에이전트는 어떤 거래도 실행하지 않으며 공개 자료만 사용합니다.", 9)

    doc.add_heading("요약", level=2)
    _table(doc, ["#", "종목명", "시장", "통화", "시총", "5일 상승률", "거래량 배수", "산업 트렌드", "등급"], [
        [i, f"{p['name']} ({p['ticker']})", p["market"], p["currency"], fmt_mcap(p["mcap"], p["currency"]),
         f"{p['ret_5d'] * 100:+.1f}%", f"{p['vol_ratio']:.1f}x", p["trend_status"], p["grade"]]
        for i, p in enumerate(picks, 1)])
    if len(picks) < 10:
        _para(doc, f"※ 필터와 검증을 모두 통과한 종목이 {len(picks)}개뿐이라 10개를 채우지 못했습니다 (임계값을 낮추지 않았습니다).", 9)

    doc.add_heading("종목별 분석", level=2)
    for i, p in enumerate(picks, 1):
        a = p["analysis"]
        doc.add_heading(f"{i}. {p['name']} ({p['ticker']})", level=3)
        crit = ", ".join(k for k, v in p["criteria"].items() if v)
        rows = [
            ("시장 · 통화 · 시총", f"{p['market']} · {p['currency']} · {fmt_mcap(p['mcap'], p['currency'])}"),
            ("5일 상승률 · 거래량", f"{p['ret_5d'] * 100:+.1f}% · 직전 5거래일 평균 대비 {p['vol_ratio']:.1f}배 "
                                 f"(종가 {p['last_close']:,.2f} {p['currency']}, {p['last_date']}) [yfinance, {p['ticker']}]"),
            ("상승 이유", mask_assertive(a["reason"])),
            ("이벤트 · 성격", f"{a['event_type']} · {a['nature']}"),
            ("산업 트렌드", p["trend_status"] + (f"\n근거: {_cites([t['article'] for t in p['trends']])}" if p["trends"] else "")),
            ("지속 가능성 근거", crit or "없음"),
            ("하방 리스크", "\n".join(f"- {mask_assertive(r)}" for r in a["risks"])),
            ("추천 등급", f"{p['grade']}  (검토 후보 · 관심 종목 · 주의 필요 중 하나이며, 매수·매도 지시가 아닙니다)"),
            ("근거 출처", _cites(a["evidence"]) + (f"\n{p['institutional']['source']} {p['institutional']['desc']}"
                                                  if p.get("institutional") and p["institutional"].get("ok") else "")),
        ]
        if p.get("track_record"):
            rows.append(("과거 추천 이력", p["track_record"]))
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for k, v in rows:
            c = t.add_row().cells
            _cell(c[0], k, True)
            _cell(c[1], v)

    doc.add_heading("검증 · 실행 정보", level=2)
    for line in meta.get("notes", []):
        _para(doc, f"• {line}", 9)
    if meta.get("excluded"):
        _para(doc, "검증 · 지속 가능성 기준 미달로 제외된 종목:", 9, True)
        for line in meta["excluded"]:
            _para(doc, f"• {line}", 9)
    _para(doc, f"토큰 사용 {meta['tokens']:,} / 60,000 · 소요 {meta['elapsed']:.0f}초 / 60초", 9)

    _disclaimer(doc)  # 마지막 단락
    doc.save(str(path))
    return path
